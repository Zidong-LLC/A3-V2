"""Ejemplos en Excel de la carga masiva (pedido del usuario, 2026-10-29): 20 filas por tipo,
se suben tal cual y el plan sale sin errores."""
import pytest

from app import imports, imports_aplicar, imports_ejemplo
from app import imports_archivo as archivo
from app.main import app

TIPOS = ("clientes", "precios", "portafolio", "motorizados")


def _leer(tipo):
    data = imports_ejemplo.ejemplo_xlsx(tipo)
    encabezados, filas = archivo.leer_tabla(data, f"ejemplo-{tipo}.xlsx")
    mapa = archivo.mapa_automatico(encabezados, tipo)
    return encabezados, filas, mapa, archivo.filas_con_mapa(filas, mapa)


@pytest.mark.parametrize("tipo", TIPOS)
def test_el_ejemplo_trae_20_filas_y_se_reconoce_solo(tipo):
    encabezados, filas, mapa, _ = _leer(tipo)
    assert len(filas) == 20
    assert archivo.detectar_tipo(encabezados, list(TIPOS)) == tipo
    assert not archivo.faltantes(mapa, tipo)
    assert set(mapa) == set(archivo.CAMPOS[tipo]), "todas las columnas del tipo se reconocen por su nombre"


def test_clientes_de_ejemplo_se_crean_todos_con_su_motorizado():
    _, _, _, filas = _leer("clientes")
    motorizados = {imports.norm(n): f"id-{n}" for n in imports_ejemplo._MOTORIZADOS_A3}
    plan = imports.plan_clientes(filas, [], lambda a, b: False, False, motorizados, {})
    assert len(plan["crear"]) == 20 and plan["errores"] == []
    assert all(c.get("courier_id") and c.get("tax_id") and c.get("email") for c in plan["crear"])
    assert len({c["tax_id"] for c in plan["crear"]}) == 20, "NIT distintos"


def test_precios_de_ejemplo_son_los_del_catalogo_y_no_cambian_nada():
    _, _, _, filas = _leer("precios")
    catalogo = [{"code": c, "name": n, "price": p} for c, n, p in imports_ejemplo.PRECIOS_REALES]
    plan = imports.plan_precios(filas, catalogo, [])
    assert plan["iguales"] == 20 and plan["actualizar"] == [] and plan["errores"] == []


def test_portafolio_de_ejemplo_crea_16_analisis_y_4_perfiles():
    _, _, _, filas = _leer("portafolio")
    plan = imports.plan_portafolio(filas, [], [])
    assert len(plan["crear"]) == 20 and plan["errores"] == []
    assert sum(c["tabla"] == "catalog_profiles" for c in plan["crear"]) == 4
    assert {c["species"] for c in plan["crear"]} <= {"canino", "felino", "ambos"}
    assert all(c["category"] and c["sample"] for c in plan["crear"])


def test_motorizados_de_ejemplo_se_crean_con_telefono():
    _, _, _, filas = _leer("motorizados")
    plan = imports_aplicar.plan_motorizados(filas, [])
    assert len(plan["crear"]) == 20 and plan["errores"] == []
    assert all(c.get("phone") for c in plan["crear"])


def test_el_libro_trae_la_hoja_como_usar_despues_de_los_datos():
    import io

    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(imports_ejemplo.ejemplo_xlsx("clientes")))
    assert libro.sheetnames == ["Clientes", "Cómo usar"]
    assert libro["Clientes"].freeze_panes == "A2"
    assert "20 filas" in str(libro["Cómo usar"]["A1"].value)


def test_la_ruta_descarga_el_excel_de_ejemplo():
    app.config["TESTING"] = True
    c = app.test_client()
    with c.session_transaction() as s:
        s["dashboard_authenticated"] = True
        s["dashboard_username"] = "admin"
    r = c.get("/cargas/ejemplo/clientes")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert 'filename="ejemplo-clientes.xlsx"' in r.headers["Content-Disposition"]
    assert c.get("/cargas/ejemplo/otra-cosa").status_code == 302
    assert "/cargas/ejemplo/clientes" in c.get("/cargas").get_data(as_text=True)
