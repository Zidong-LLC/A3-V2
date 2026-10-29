"""Mejoras de la carga masiva (pedido del usuario, 2026-10-29): archivos con título arriba,
datos en la segunda hoja, sedes con el mismo NIT, zona y factura electrónica, plantillas en
Excel con desplegables y exportación del catálogo completo."""
import io
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app import imports, imports_exportar, imports_plantillas
from app import imports_archivo as archivo
from app.main import app

TIPOS = ["precios", "clientes", "portafolio", "motorizados"]


def _xlsx(hojas) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas:
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _filas(csv_texto: str, tipo: str) -> list[dict]:
    enc, filas = archivo.leer_tabla(csv_texto.encode("utf-8"), "x.csv")
    return archivo.filas_con_mapa(filas, archivo.mapa_automatico(enc, tipo))


# ── Lectura ─────────────────────────────────────────────────────────────────────────


def test_un_titulo_y_filas_vacias_arriba_del_encabezado_no_lo_tapan():
    data = _xlsx([("Hoja1", [["LISTA DE CLIENTES 2026"], [], ["Nombre del cliente", "N.I.T.", "Celular", "Localidad"],
                             ["Vet Norte", "900.123.456-1", 3001234567, "Suba"]])])
    enc, filas = archivo.leer_tabla(data, "clientes.xlsx")
    assert enc == ["Nombre del cliente", "N.I.T.", "Celular", "Localidad"] and len(filas) == 1
    assert archivo.detectar_tipo(enc, TIPOS) == "clientes"
    assert archivo.mapa_automatico(enc, "clientes") == {"name": 0, "tax_id": 1, "phone": 2, "zone": 3}


def test_se_lee_la_hoja_que_tiene_los_datos_aunque_no_sea_la_primera():
    data = _xlsx([("Portada", [["A3 Laboratorio"], ["Precios vigentes"]]),
                  ("Precios", [["Código", "Examen", "Precio Venta"], [1101, "Cuadro Hemático", "$ 14.000"]])])
    enc, filas = archivo.leer_tabla(data, "precios.xlsx")
    assert enc == ["Código", "Examen", "Precio Venta"] and filas == [["1101", "Cuadro Hemático", "$ 14.000"]]
    assert archivo.detectar_tipo(enc, TIPOS) == "precios"


def test_un_csv_sin_titulo_se_lee_igual_que_antes():
    enc, filas = archivo.leer_tabla(b"codigo;precio\n952;35000\n", "p.csv")
    assert enc == ["codigo", "precio"] and filas == [["952", "35000"]]


def test_un_csv_sin_ningun_encabezado_reconocible_usa_la_primera_fila():
    enc, filas = archivo.leer_tabla(b"a;b\n1;2\n", "p.csv")
    assert enc == ["a", "b"] and filas == [["1", "2"]]


# ── Sedes, zona y factura electrónica ───────────────────────────────────────────────

SEDES = [{"id": "s1", "clinic_name": "Club Animals Norte", "tax_id": "900111", "email": None, "address": "Cl 1", "zone": None},
         {"id": "s2", "clinic_name": "Club Animals Sur", "tax_id": "900111", "email": None, "address": "Cl 2", "zone": None}]


def _coincide(a, b):
    return imports.norm(a) == imports.norm(b)


def test_una_sede_nombrada_se_completa_y_otra_sede_del_mismo_nit_no_es_repetida():
    filas = _filas("nombre,nit,zona\nClub Animals Norte,900111,Usaquén\nClub Animals Sur,900111,Bosa\n", "clientes")
    plan = imports.plan_clientes(filas, SEDES, _coincide)
    assert plan["errores"] == []
    assert [(c["id"], c["cambios"]) for c in plan["actualizar"]] == [("s1", {"zone": "Usaquén"}), ("s2", {"zone": "Bosa"})]


def test_sin_sede_reconocible_solo_se_completa_el_correo_como_antes():
    filas = _filas("nombre,nit,correo,direccion\nClub Animals,900111,f@club.co,Cl 9\n", "clientes")
    plan = imports.plan_clientes(filas, SEDES, lambda a, b: False)
    assert plan["crear"] == [] and plan["errores"] == []
    assert [c["cambios"] for c in plan["actualizar"]] == [{"email": "f@club.co"}, {"email": "f@club.co"}]


def test_zona_y_factura_electronica_se_crean_y_se_actualizan():
    filas = _filas("nombre,nit,zona,factura electronica\nVet Nueva,901000,Kennedy,si\nVet Vieja,902000,Suba,no\n", "clientes")
    vieja = {"id": "v1", "clinic_name": "Vet Vieja", "tax_id": "902000", "zone": "", "electronic_invoice": True}
    plan = imports.plan_clientes(filas, [vieja], _coincide, actualizar=True)
    assert plan["crear"] == [{"clinic_name": "Vet Nueva", "tax_id": "901000", "zone": "Kennedy", "electronic_invoice": True}]
    assert plan["actualizar"][0]["cambios"] == {"zone": "Suba", "electronic_invoice": False}


def test_una_celda_de_factura_que_no_se_entiende_se_ignora():
    filas = _filas("nombre,nit,factura electronica\nVet X,903000,tal vez\n", "clientes")
    plan = imports.plan_clientes(filas, [], _coincide)
    assert plan["crear"] == [{"clinic_name": "Vet X", "tax_id": "903000"}]


def test_la_exportacion_de_clientes_trae_zona_y_factura_y_vuelve_a_entrar():
    cliente = {"id": "c1", "clinic_name": "Vet Z", "tax_id": "904000", "phone": "3001", "address": "Cl 1", "zone": "Bosa",
               "city": "Bogotá", "email": "z@z.co", "electronic_invoice": True, "is_active": True}
    with patch("app.imports_exportar.db.list_couriers", return_value=[]), \
            patch("app.imports_exportar.maestros_db.asignaciones", return_value={}), \
            patch("app.imports_exportar.db.list_clients_basic", return_value=[cliente]):
        texto = imports_exportar.exportar("clientes")
    enc, filas = archivo.leer_tabla(texto.encode("utf-8"), "c.csv")
    assert enc[:8] == ["nombre", "nit", "telefono", "direccion", "zona", "municipio", "correo", "factura electronica"]
    plan = imports.plan_clientes(archivo.filas_con_mapa(filas, archivo.mapa_automatico(enc, "clientes")), [cliente], _coincide)
    assert plan["iguales"] == 1 and plan["actualizar"] == [] and plan["crear"] == []


def test_la_exportacion_del_catalogo_completo_tiene_las_columnas_del_portafolio():
    with patch("app.imports_exportar.db.list_catalog_tests", return_value=[{"code": "1101", "name": "Hemograma", "price": 14000, "category": "Hematología", "species": "ambos", "sample": "EDTA"}]), \
            patch("app.imports_exportar.db.list_catalog_profiles", return_value=[{"code": "101", "name": "Perfil I", "price": 30000, "category": "Perfiles", "species": "canino"}]):
        texto = imports_exportar.exportar("portafolio")
    enc, filas = archivo.leer_tabla(texto.encode("utf-8"), "p.csv")
    assert enc == ["codigo", "nombre", "precio", "tipo", "categoria", "especie", "muestra"]
    assert [f[3] for f in filas] == ["analisis", "perfil"]
    # Con categoría y muestra se reconoce como portafolio: los códigos existentes quedan «sin
    # cambios» y los nuevos se crean. Para cambiar precios se elige «Precios» en el desplegable.
    assert archivo.detectar_tipo(enc, TIPOS) == "portafolio"


# ── Pantalla: ambigüedad y plantillas ───────────────────────────────────────────────


def _client():
    app.config["TESTING"] = True
    c = app.test_client()
    with c.session_transaction() as s:
        s["dashboard_authenticated"] = True
        s["dashboard_username"] = "admin"
    return c


def test_solo_nombre_y_telefono_en_auto_pregunta_en_vez_de_adivinar():
    c = _client()
    r = c.post("/cargas/previsualizar", data={"tipo": "auto", "csv": (io.BytesIO(b"nombre,telefono\nVet Sur,3109876543\n"), "l.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "elige si son clientes o motorizados" in r.get_data(as_text=True)


def test_la_plantilla_es_excel_con_desplegables_y_dos_filas_de_ejemplo():
    data = imports_plantillas.plantilla_xlsx("clientes", ["Javier", "Luis"])
    libro = load_workbook(io.BytesIO(data))
    hoja = libro["Clientes"]
    assert [c.value for c in hoja[1]] == ["nombre", "nit", "telefono", "direccion", "zona", "municipio", "correo", "factura electronica", "motorizado"]
    assert hoja.max_row == 3 and libro.sheetnames == ["Clientes", "Cómo usar"]
    formulas = [dv.formula1 for dv in hoja.data_validations.dataValidation]
    assert '"si,no"' in formulas and '"Javier,Luis"' in formulas
    enc, filas = archivo.leer_tabla(data, "plantilla.xlsx")
    assert archivo.detectar_tipo(enc, TIPOS) == "clientes" and len(filas) == 2


def test_la_plantilla_de_portafolio_tiene_desplegables_de_especie_y_tipo():
    libro = load_workbook(io.BytesIO(imports_plantillas.plantilla_xlsx("portafolio")))
    formulas = [dv.formula1 for dv in libro["Portafolio"].data_validations.dataValidation]
    assert '"canino,felino,ambos"' in formulas and '"analisis,perfil"' in formulas


def test_las_rutas_de_plantilla_devuelven_excel_y_csv_a_pedido():
    c = _client()
    with patch("app.dashboard_import.db.list_active_couriers", return_value=[{"name": "Javier"}]):
        r = c.get("/cargas/plantilla/clientes")
    assert r.status_code == 200 and r.data[:2] == b"PK" and 'plantilla-clientes.xlsx' in r.headers["Content-Disposition"]
    r = c.get("/cargas/plantilla/precios?formato=csv")
    assert r.status_code == 200 and r.get_data(as_text=True).startswith("﻿codigo;precio")
