"""Plantillas en Excel de las cargas masivas (pedido del usuario, 2026-10-29).

La plantilla CSV de una fila no ayudaba a nadie: no decía qué valores acepta cada columna.
Estas traen los encabezados que la carga reconoce, dos filas de ejemplo en gris (se borran
y se pisan con los datos reales), desplegables donde hay valores fijos (especie, tipo, sí/no,
motorizados reales de A3) y la misma hoja «Cómo usar» que los ejemplos.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.imports_ejemplo import EJEMPLOS, _hoja_como_usar, _estilizar

_FILAS_EDITABLES = 500
# Columna → lista de valores permitidos (desplegable). Los motorizados se pasan por parámetro.
_LISTAS = {
    "especie": ["canino", "felino", "ambos"],
    "tipo": ["analisis", "perfil"],
    "factura electronica": ["si", "no"],
}


def plantilla_xlsx(tipo: str, motorizados: list[str] | None = None) -> bytes:
    encabezados, filas = EJEMPLOS[tipo]
    libro = Workbook()
    hoja = libro.active
    hoja.title = tipo.capitalize()
    hoja.append(encabezados)
    ejemplo = Font(italic=True, color="9A9A9A")
    for fila in filas()[:2]:
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            celda.font = ejemplo
    _estilizar(hoja, encabezados)
    for indice, nombre in enumerate(encabezados, start=1):
        valores = list(_LISTAS.get(nombre, []))
        if nombre == "motorizado":
            valores = [m for m in (motorizados or []) if m]
        if not valores:
            continue
        lista = DataValidation(type="list", formula1='"' + ",".join(valores)[:250] + '"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Valor no válido",
                               error=f"Elige uno de: {', '.join(valores)}")
        hoja.add_data_validation(lista)
        letra = get_column_letter(indice)
        lista.add(f"{letra}2:{letra}{_FILAS_EDITABLES}")
    _hoja_como_usar(libro, tipo, "Plantilla: borra las dos filas grises de ejemplo y pega tus datos debajo "
                                 "del encabezado. Las columnas con flecha tienen desplegable.")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
