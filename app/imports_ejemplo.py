"""Ejemplos en Excel de las cargas masivas (pedido del usuario, 2026-10-29).

Las plantillas CSV traen una sola fila y no muestran cómo se ve un archivo real. Acá va un
.xlsx por tipo con 20 filas verosímiles, encabezado con los colores de A3, anchos de columna
y una segunda hoja «Cómo usar». La carga lee la PRIMERA hoja, con los mismos encabezados que
las plantillas, así que el ejemplo se puede subir tal cual para ver el plan.

Los datos son inventados (clientes, motorizados y portafolio) salvo los precios, que usan
códigos y valores reales del catálogo: subirlos da «sin cambios» y no toca nada.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_VINO = "7A0D20"
_ZEBRA = "FBF3F4"

_MOTORIZADOS_A3 = ("Alexander", "Cesar", "Diego", "Gerardo", "Javier", "Jeeferson", "Luis", "Marlon")
_LOCALIDADES = ("Kennedy", "Suba", "Engativá", "Usaquén", "Chapinero", "Fontibón", "Bosa",
                "Teusaquillo", "Puente Aranda", "Ciudad Bolívar", "Soacha", "Chía", "Cajicá",
                "Mosquera", "Funza", "Rafael Uribe Uribe", "Tunjuelito", "Barrios Unidos", "Usme", "La Calera")
_NOMBRES = ("Veterinaria Los Alpes", "Clínica Veterinaria Patitas Felices", "Centro Veterinario San Martín",
            "Veterinaria El Bosque", "Clínica Animal Care Suba", "Veterinaria Huellitas del Norte",
            "Hospital Veterinario La Sabana", "Veterinaria Zoovet Kennedy", "Clínica Veterinaria Canes y Gatos",
            "Veterinaria Mi Mascota Fontibón", "Dra. Laura Gómez — Consultorio", "Veterinaria Amigos Fieles",
            "Clínica Veterinaria Santa Bárbara", "Veterinaria Pet Center Bosa", "Centro Médico Veterinario Vida",
            "Veterinaria Los Pinos", "Clínica Veterinaria Dog House", "Veterinaria Animalia Chía",
            "Dr. Andrés Rojas — Consultorio", "Veterinaria Colitas Alegres")
_CALLES = ("Cl 45 Sur # 78-12", "Cra 91 # 139-20", "Cl 80 # 70-15", "Cra 7 # 116-35", "Cl 63 # 11-40",
           "Cra 100 # 23-18", "Cl 59 Sur # 80-25", "Cra 24 # 39-16", "Cl 13 # 50-22", "Cra 18 Bis # 60 Sur-14",
           "Cl 3 # 10-45", "Cra 9 # 12-30", "Cl 125 # 19-60", "Cra 80 # 65 Sur-10", "Cl 72 # 20-33",
           "Cra 50 # 22-17", "Cl 134 # 55-28", "Cra 11 # 15-70", "Cl 100 # 8-25", "Cra 68 # 35-40")

PRECIOS_REALES = [
    ("1101", "Cuadro Hemático Completo", 14000), ("1102", "Prueba Cruzada de Coombs", 28000),
    ("1103", "Recuento de Plaquetas", 7000), ("1104", "Recuento de Reticulocitos", 8000),
    ("1105", "Hemoparásitos", 10000), ("1106", "Hemoglobina y Hematocrito", 8000),
    ("1107", "Células L.E. (Manual)", 22000), ("1108", "Hemocitología (Registro Fotográfico)", 18000),
    ("1109", "Prueba de Coombs", 28000), ("1201", "PT (Tiempo de Protrombina)", 18000),
    ("1202", "PTT (Tiempo parcial de Tromboplastina)", 18000), ("1203", "Dímero D - Vcheck", 65000),
    ("1701", "Coprológico", 12000), ("1702", "Coproscópico", 15000),
    ("2057", "Panleucopenia Felina (Antígeno)", 55000), ("2067", "Panleucopenia Vcheck (Antígeno)", 63000),
    ("101", "Perfil Parasitológico I", 30000), ("102", "Perfil Parasitológico II", 23000),
    ("103", "Perfil Parasitológico III", 40000), ("104", "Perfil Parasitológico IV", 25000),
]

_PORTAFOLIO = [
    ("9901", "Ácido Fólico", 48000, "analisis", "Química", "ambos", "Tubo Rojo o Amarillo"),
    ("9902", "Vitamina B12", 52000, "analisis", "Química", "ambos", "Tubo Rojo o Amarillo"),
    ("9903", "Lipasa Pancreática Felina (fPL)", 95000, "analisis", "Inmunológico", "felino", "Tubo Rojo o Amarillo"),
    ("9904", "Lipasa Pancreática Canina (cPL)", 95000, "analisis", "Inmunológico", "canino", "Tubo Rojo o Amarillo"),
    ("9905", "Fructosamina", 38000, "analisis", "Química", "ambos", "Tubo Rojo o Amarillo"),
    ("9906", "Ácidos Biliares Pre y Post", 70000, "analisis", "Química", "ambos", "Tubo Rojo o Amarillo"),
    ("9907", "Cortisol Basal", 45000, "analisis", "Hormonas", "ambos", "Tubo Rojo o Amarillo"),
    ("9908", "Progesterona", 48000, "analisis", "Hormonas", "canino", "Tubo Rojo o Amarillo"),
    ("9909", "Urocultivo con Antibiograma", 42000, "analisis", "Microbiología", "ambos", "Orina estéril"),
    ("9910", "Cultivo de Piel y Antibiograma", 42000, "analisis", "Microbiología", "ambos", "Hisopo estéril"),
    ("9911", "Citología de Oído", 20000, "analisis", "Citología", "ambos", "2 láminas"),
    ("9912", "Citología de Masa", 35000, "analisis", "Citología", "ambos", "2 láminas"),
    ("9913", "Giardia (Antígeno)", 38000, "analisis", "Inmunológico", "ambos", "Materia Fecal"),
    ("9914", "Parvovirus (Antígeno)", 38000, "analisis", "Inmunológico", "canino", "Materia Fecal"),
    ("9915", "Fibrinógeno", 18000, "analisis", "Coagulación", "ambos", "Tubo Tapa Azul con 3/4 de sangre"),
    ("9916", "Relación Proteína/Creatinina en Orina", 30000, "analisis", "Uroanálisis", "ambos", "Orina"),
    ("9917", "Perfil Geriátrico Canino", 120000, "perfil", "Perfiles", "canino", "Tubo Rojo o Amarillo y Tapa Morada"),
    ("9918", "Perfil Geriátrico Felino", 125000, "perfil", "Perfiles", "felino", "Tubo Rojo o Amarillo y Tapa Morada"),
    ("9919", "Perfil Hepático Completo", 85000, "perfil", "Perfiles", "ambos", "Tubo Rojo o Amarillo"),
    ("9920", "Perfil Tiroideo Felino", 90000, "perfil", "Perfiles", "felino", "Tubo Rojo o Amarillo"),
]

_MOTORIZADOS = ("Camilo Andrés Torres", "Juan Pablo Ríos", "Oscar Mauricio Pérez", "Wilson Castañeda",
                "Édgar Mora", "Fabián Suárez", "Jhon Freddy Molina", "Néstor Guzmán", "Yeison Pardo",
                "Carlos Alberto Niño", "Hernán Darío Ortiz", "Miguel Ángel Cruz", "Sebastián Lozano",
                "Andrés Felipe Vargas", "Jorge Luis Peña", "Cristian Camilo Díaz", "Brayan Rodríguez",
                "Daniel Esteban Ruiz", "Julián Cárdenas", "Mauricio Hernández")


def _nit(i: int) -> str:
    base = 901_200_000 + i * 7_919
    digito = (base * 3) % 10
    return f"{base}-{digito}"


def _telefono(i: int) -> str:
    return f"31{i % 10}{(5_200_000 + i * 48_611) % 10_000_000:07d}"


def _correo(nombre: str, i: int) -> str:
    palabra = [p for p in nombre.lower().replace("—", " ").split() if p not in ("veterinaria", "clínica", "centro", "dra.", "dr.", "consultorio", "hospital")]
    return f"facturacion{i}@{(palabra[0] if palabra else 'clinica').replace('.', '')}.co"


_MUNICIPIOS = {"Soacha", "Chía", "Cajicá", "Mosquera", "Funza", "La Calera"}


def _clientes() -> list[list]:
    filas = []
    for i, nombre in enumerate(_NOMBRES, start=1):
        lugar = _LOCALIDADES[i - 1]
        zona, municipio = ("", lugar) if lugar in _MUNICIPIOS else (lugar, "Bogotá")
        filas.append([nombre, _nit(i), _telefono(i), _CALLES[i - 1], zona, municipio,
                      _correo(nombre, i), "si" if i % 3 else "no", _MOTORIZADOS_A3[i % len(_MOTORIZADOS_A3)]])
    return filas


EJEMPLOS = {
    "clientes": (["nombre", "nit", "telefono", "direccion", "zona", "municipio", "correo", "factura electronica", "motorizado"],
                 _clientes),
    "precios": (["codigo", "nombre", "precio"], lambda: [list(f) for f in PRECIOS_REALES]),
    "portafolio": (["codigo", "nombre", "precio", "tipo", "categoria", "especie", "muestra"],
                   lambda: [list(f) for f in _PORTAFOLIO]),
    "motorizados": (["nombre", "telefono"], lambda: [[n, _telefono(i + 40)] for i, n in enumerate(_MOTORIZADOS)]),
}

_COMO_USAR = {
    "clientes": ("Un cliente por fila. Con el NIT se cruza contra la base: si ya existe, solo se completan los datos "
                 "vacíos (o se actualizan, si marcas la casilla). Sedes con el mismo NIT: se reconocen por el nombre. "
                 "Zona = localidad de Bogotá (con ella el bot asigna motorizado); municipio para los de afuera. "
                 "Factura electrónica: si/no. El motorizado se asigna por su nombre tal como está en la plataforma."),
    "precios": ("Un código por fila con su precio nuevo. Los códigos tienen que existir en el catálogo; la carga "
                "muestra antes/después y solo cambia los que difieren. Este ejemplo trae los precios actuales: subirlo "
                "da «sin cambios»."),
    "portafolio": ("Análisis o perfiles NUEVOS (códigos que todavía no están en el catálogo). Tipo: analisis o perfil. "
                   "Especie: canino, felino o ambos (vacío = ambos). Categoría obligatoria. Los códigos 99xx de este "
                   "ejemplo son inventados."),
    "motorizados": ("Un motorizado por fila: nombre y celular. Si el nombre ya existe se actualiza el teléfono; si no, "
                    "se crea activo."),
}
_PASOS = ("1. Deja esta hoja como está o reemplaza las filas por tus datos (los encabezados no se cambian).",
          "2. En la plataforma: Herramientas → Cargas por archivo → elige el archivo → «Ver qué va a pasar».",
          "3. Revisa el plan (qué se crea, qué se actualiza, qué queda igual) y recién ahí «Aplicar estos cambios».",
          "Acepta también CSV con coma o punto y coma; las columnas se reconocen por su nombre y en cualquier orden.",
          "Una lista SIN encabezado (el código y al lado el precio desde la primera fila) también se reconoce; "
          "y si el código viene pegado al nombre («1101 - Cuadro Hemático»), se separa solo.")


def _estilizar(hoja, encabezados: list[str]) -> None:
    """Encabezado con los colores de A3, filas alternadas, precios con formato, anchos,
    encabezado fijo y filtro. Lo comparten el ejemplo y la plantilla."""
    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=_VINO)
    for celda in hoja[1]:
        celda.font, celda.fill = cabecera, relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")
    zebra = PatternFill("solid", fgColor=_ZEBRA)
    for numero, fila in enumerate(hoja.iter_rows(min_row=2), start=2):
        for celda in fila:
            if numero % 2 == 0:
                celda.fill = zebra
            if encabezados[celda.column - 1] == "precio":
                celda.number_format = '"$"#,##0'
    for indice, nombre in enumerate(encabezados, start=1):
        ancho = max(len(str(c.value or "")) for c in hoja[get_column_letter(indice)])
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 3, 12), 46)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions


def _hoja_como_usar(libro, tipo: str, titulo: str) -> None:
    ayuda = libro.create_sheet("Cómo usar")
    ayuda.column_dimensions["A"].width = 110
    celda = ayuda.cell(row=1, column=1, value=f"Carga masiva · {tipo} · {titulo}")
    celda.font = Font(bold=True, size=13, color=_VINO)
    ayuda.cell(row=3, column=1, value=_COMO_USAR[tipo]).alignment = Alignment(wrap_text=True, vertical="top")
    ayuda.row_dimensions[3].height = 60
    for n, paso in enumerate(_PASOS, start=5):
        ayuda.cell(row=n, column=1, value=paso).alignment = Alignment(wrap_text=True)


def ejemplo_xlsx(tipo: str) -> bytes:
    """El libro de ejemplo del tipo: hoja de datos (la que lee la carga) + hoja «Cómo usar»."""
    encabezados, filas = EJEMPLOS[tipo]
    libro = Workbook()
    hoja = libro.active
    hoja.title = tipo.capitalize()
    hoja.append(encabezados)
    for fila in filas():
        hoja.append(fila)
    _estilizar(hoja, encabezados)
    _hoja_como_usar(libro, tipo, f"ejemplo con {len(filas())} filas")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
